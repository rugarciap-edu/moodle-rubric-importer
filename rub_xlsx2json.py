#!/usr/bin/env python3
"""Convierte una rúbrica Excel (plantilla_rubrica.xlsx) a JSON (formato_version 1).

Uso:
    python rub_xlsx2json.py rubrica.xlsx [salida.json]
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

FILA_INICIO = 5
BOOLEANOS = {"sí": True, "si": True, "no": False}


def vacio(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def num(v, donde: str):
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        try:
            v = float(str(v).replace(",", "."))
        except ValueError:
            raise SystemExit(f"{donde}: «{v}» no es un número.")
    return int(v) if float(v).is_integer() else v


def leer_general(ws) -> dict[str, object]:
    valores = {}
    for fila in range(4, ws.max_row + 1):
        clave = ws.cell(fila, 1).value
        if clave and not clave.startswith("_"):
            valores[clave] = ws.cell(fila, 3).value
    return valores


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("salida", type=Path, nargs="?")
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, data_only=True)
    g = leer_general(wb["General"])

    if vacio(g.get("nombre")):
        raise SystemExit("General: falta el nombre de la rúbrica.")
    n = g.get("escala.numero_niveles")
    if vacio(n):
        raise SystemExit("General: falta el número de niveles.")
    n = num(n, "escala.numero_niveles")
    nombres = [g.get(f"escala.nombres_niveles.{i}") for i in range(1, n + 1)]
    if any(vacio(x) for x in nombres):
        raise SystemExit("General: faltan nombres de niveles en la escala.")
    nombres = [str(x).strip() for x in nombres]

    rubrica: dict = {"formato_version": 1, "nombre": str(g["nombre"]).strip()}
    if not vacio(g.get("descripcion")):
        rubrica["descripcion"] = str(g["descripcion"]).strip()
    if not vacio(g.get("puntuacion_maxima")):
        rubrica["puntuacion_maxima"] = num(g["puntuacion_maxima"], "puntuacion_maxima")
    rubrica["escala"] = {"numero_niveles": n, "nombres_niveles": nombres}
    rubrica["moodle"] = {
        "url_tarea": None if vacio(g.get("moodle.url_tarea")) else str(g["moodle.url_tarea"]).strip()
    }

    opciones = {}
    for clave, v in g.items():
        if clave.startswith("opciones.") and not vacio(v):
            texto = str(v).strip().lower()
            if texto not in BOOLEANOS:
                raise SystemExit(f"General: {clave} debe ser «Sí» o «No» (valor: {v}).")
            opciones[clave.split(".", 1)[1]] = BOOLEANOS[texto]
    if opciones:
        rubrica["opciones"] = opciones

    # Agrupa filas por «Criterio nº» respetando el orden de aparición
    ws = wb["Criterios"]
    criterios: dict[object, dict] = {}
    for fila in range(FILA_INICIO, ws.max_row + 1):
        a, b, c, d, e, f = (ws.cell(fila, k).value for k in range(1, 7))
        if all(vacio(x) for x in (a, b, c, d, e, f)):
            continue
        donde = f"Criterios fila {fila}"
        if vacio(a):
            raise SystemExit(f"{donde}: falta «Criterio nº».")
        crit = criterios.setdefault(num(a, donde), {"descripcion": "", "niveles": []})
        if not vacio(b) and not crit["descripcion"]:
            crit["descripcion"] = str(b).strip()
        if vacio(c) or vacio(e) or vacio(f):
            raise SystemExit(f"{donde}: faltan «Nivel nº», «Puntos» o «Definición».")
        idx = num(c, donde)
        if not isinstance(idx, int) or not 1 <= idx <= n:
            raise SystemExit(f"{donde}: «Nivel nº» debe estar entre 1 y {n}.")
        crit["niveles"].append(
            (idx, {"nombre": nombres[idx - 1], "puntos": num(e, donde), "definicion": str(f).strip()})
        )

    if not criterios:
        raise SystemExit("Criterios: no hay ningún criterio.")
    salida_crit = []
    for cid, crit in criterios.items():
        if not crit["descripcion"]:
            raise SystemExit(f"Criterio {cid}: falta la descripción.")
        indices = [i for i, _ in crit["niveles"]]
        if len(set(indices)) != len(indices) or len(indices) != n:
            raise SystemExit(f"Criterio {cid}: debe tener exactamente un nivel por cada uno de los {n}.")
        niveles = [niv for _, niv in sorted(crit["niveles"], key=lambda t: t[0])]
        puntos = [x["puntos"] for x in niveles]
        if len(set(puntos)) != len(puntos):
            raise SystemExit(f"Criterio {cid}: los puntos no pueden repetirse.")
        salida_crit.append({"descripcion": crit["descripcion"], "niveles": niveles})
    rubrica["criterios"] = salida_crit

    maximo = sum(max(x["puntos"] for x in c["niveles"]) for c in salida_crit)
    if "puntuacion_maxima" in rubrica and rubrica["puntuacion_maxima"] != maximo:
        print(
            f"Aviso: puntuacion_maxima ({rubrica['puntuacion_maxima']}) "
            f"≠ suma de máximos ({maximo}).",
            file=sys.stderr,
        )

    salida = args.salida or args.xlsx.with_suffix(".json")
    salida.write_text(json.dumps(rubrica, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Escrito {salida} ({len(salida_crit)} criterios, {n} niveles).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
