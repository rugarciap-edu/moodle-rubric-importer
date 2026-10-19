#!/usr/bin/env python3
"""Convierte una rúbrica JSON (formato_version 1) a Excel usando la plantilla.

Uso:
    python rub_json2xlsx.py rubrica.json [salida.xlsx] [--plantilla RUTA]
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

PLANTILLA = Path(__file__).parent / "templates" / "plantilla_rubrica.xlsx"
FILA_INICIO = 5
FILA_FIN = 304  # capacidad de la hoja Criterios (300 filas)
FILA_NOMBRES_ESCALA = 14
MAX_NIVELES = 10
OPCIONES_FILAS = range(25, 34)


def num(valor):
    return int(valor) if isinstance(valor, float) and valor.is_integer() else valor


def construir_escala(datos: dict) -> tuple[int, list[str]]:
    criterios = datos["criterios"]
    esc = datos.get("escala")
    if esc:
        return esc["numero_niveles"], list(esc["nombres_niveles"])
    n = len(criterios[0]["niveles"])
    if any(len(c["niveles"]) != n for c in criterios):
        raise SystemExit(
            "La plantilla Excel exige el mismo número de niveles en todos los criterios."
        )
    nombres = [niv.get("nombre", "") for niv in criterios[0]["niveles"]]
    return n, nombres


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("json", type=Path)
    ap.add_argument("salida", type=Path, nargs="?")
    ap.add_argument("--plantilla", type=Path, default=PLANTILLA)
    args = ap.parse_args()

    datos = json.loads(args.json.read_text(encoding="utf-8"))
    if datos.get("formato_version") != 1:
        raise SystemExit("formato_version no soportado (se espera 1).")
    criterios = datos["criterios"]
    n_niveles, nombres = construir_escala(datos)
    if not 2 <= n_niveles <= MAX_NIVELES or len(nombres) != n_niveles:
        raise SystemExit("Escala inválida: entre 2 y 10 niveles con un nombre por nivel.")
    if any(len(c["niveles"]) != n_niveles for c in criterios):
        raise SystemExit("Todos los criterios deben tener tantos niveles como la escala.")
    if len(criterios) * n_niveles > FILA_FIN - FILA_INICIO + 1:
        raise SystemExit("Demasiados criterios/niveles para la capacidad de la plantilla.")

    wb = load_workbook(args.plantilla)
    gen, cri = wb["General"], wb["Criterios"]

    valores = {
        "nombre": datos["nombre"],
        "descripcion": datos.get("descripcion"),
        "puntuacion_maxima": datos.get("puntuacion_maxima"),
        "moodle.url_tarea": (datos.get("moodle") or {}).get("url_tarea"),
        "escala.numero_niveles": n_niveles,
    }
    for i in range(MAX_NIVELES):
        valores[f"escala.nombres_niveles.{i + 1}"] = nombres[i] if i < n_niveles else None
    opciones = datos.get("opciones") or {}
    for clave, v in opciones.items():
        valores[f"opciones.{clave}"] = "Sí" if v else "No"

    claves_usadas = set()
    for fila in range(4, gen.max_row + 1):
        clave = gen.cell(fila, 1).value
        if not clave or clave.startswith("_") or clave == "formato_version":
            continue
        if clave.startswith("opciones.") or clave in valores:
            gen.cell(fila, 3).value = num(valores.get(clave))
            claves_usadas.add(clave)
    desconocidas = {f"opciones.{k}" for k in opciones} - claves_usadas
    if desconocidas:
        raise SystemExit(f"Opciones sin columna en la plantilla: {sorted(desconocidas)}")

    for fila in range(FILA_INICIO, FILA_FIN + 1):
        for col in range(1, 7):
            cri.cell(fila, col).value = None
    fila = FILA_INICIO
    for i, crit in enumerate(criterios, 1):
        # La plantilla ordena por «Nivel nº» según los puntos crecientes
        for j, niv in enumerate(sorted(crit["niveles"], key=lambda x: x["puntos"]), 1):
            cri.cell(fila, 1).value = i
            if j == 1:
                cri.cell(fila, 2).value = crit["descripcion"]
            cri.cell(fila, 3).value = j
            cri.cell(fila, 4).value = nombres[j - 1]
            cri.cell(fila, 5).value = num(niv["puntos"])
            cri.cell(fila, 6).value = niv["definicion"]
            fila += 1

    salida = args.salida or args.json.with_suffix(".xlsx")
    wb.save(salida)
    print(f"Escrito {salida} ({len(criterios)} criterios, {n_niveles} niveles).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
